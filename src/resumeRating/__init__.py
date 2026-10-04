import argparse
import csv
import json
import os
import sys
from pathlib import Path
from typing import Any

from docx import Document
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader


MODEL = "openai/gpt-oss-120b"
DEFAULT_RESUMES_DIR = Path(__file__).resolve().parents[2] / "resumes"


class JobDescription(BaseModel):
    role: str = "Not specified"
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: float | None = None
    education_requirements: list[str] = Field(default_factory=list)
    responsibilities: list[str] = Field(default_factory=list)


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


class ResumeInformation(BaseModel):
    name: bool
    email: bool
    phone: bool
    skills: bool
    experience: bool
    education: bool
    projects: bool
    certifications: bool


class MatchAnalysis(BaseModel):
    matching_required_skills: list[str] = Field(default_factory=list)
    missing_required_skills: list[str] = Field(default_factory=list)
    matching_preferred_skills: list[str] = Field(default_factory=list)
    missing_preferred_skills: list[str] = Field(default_factory=list)
    experience_requirement_met: bool | None = None
    education_requirement_satisfied: bool | None = None
    relevant_experience: list[str] = Field(default_factory=list)
    strengths: list[str] = Field(default_factory=list)
    verdict: str = ""
    quality_issues: list[str] = Field(default_factory=list)
    improvement_suggestions: list[str] = Field(default_factory=list)


class ScoreBreakdown(BaseModel):
    required_skills: float
    preferred_skills: float
    experience: float
    education: float
    relevant_experience: float


class CandidateResult(BaseModel):
    candidate_name: str
    final_score: float
    score_breakdown: ScoreBreakdown
    matching_required_skills: list[str]
    missing_required_skills: list[str]
    matching_preferred_skills: list[str]
    missing_preferred_skills: list[str]
    experience_requirement_met: bool | None
    education_requirement_satisfied: bool | None
    relevant_experience: list[str]
    strengths: list[str]
    improvement_suggestions: list[str]
    resume_information: ResumeInformation
    resume_quality_issues: list[str]
    verdict: str


def _request_json(client: Groq, system_prompt: str, user_prompt: str) -> dict[str, Any]:
    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        response_format={"type": "json_object"},
    )
    raw_output = response.choices[0].message.content
    if not raw_output:
        raise ValueError("The LLM returned an empty response.")
    try:
        data = json.loads(raw_output)
    except json.JSONDecodeError as error:
        raise ValueError(f"The LLM returned invalid JSON: {error}") from error
    if not isinstance(data, dict):
        raise ValueError("The LLM response must be a JSON object.")
    return data


def parse_job_description(text: str, client: Groq) -> JobDescription:
    schema = json.dumps(JobDescription.model_json_schema(), indent=2)
    system_prompt = f"""Extract job requirements from the supplied job description.
Return only a JSON object matching this schema:
{schema}

Rules:
- Use only information stated in the job description; do not invent requirements.
- Keep required and preferred skills separate. Use concise skill names.
- Do not put education or years-of-experience requirements in skill lists.
- Set minimum_experience to a number only when a minimum is explicit; otherwise null.
- Return empty arrays when information is absent.
"""
    data = _request_json(client, system_prompt, text)
    try:
        return JobDescription.model_validate(data)
    except Exception as error:
        raise ValueError(f"Could not validate the parsed job description: {error}") from error


def parse_resume(text: str, client: Groq) -> Resume:
    schema = json.dumps(Resume.model_json_schema(), indent=2)
    system_prompt = f"""Extract structured information from the supplied resume.
Return only a JSON object matching this schema:
{schema}

Use only information in the resume. Do not infer or exaggerate qualifications.
Use null for absent scalar fields and empty arrays for absent list fields.
Treat internships as experience when they describe technical or professional work.
Keep projects and coursework separate from employment experience.
"""
    data = _request_json(client, system_prompt, text)
    try:
        return Resume.model_validate(data)
    except Exception as error:
        raise ValueError(f"Could not validate the parsed resume: {error}") from error


def analyze_match(job: JobDescription, resume: Resume, client: Groq) -> MatchAnalysis:
    schema = json.dumps(MatchAnalysis.model_json_schema(), indent=2)
    system_prompt = f"""Compare the job requirements with the resume and return only JSON
matching this schema:
{schema}

Rules:
- Match skills semantically, not only by identical wording. For example, OOP and
  object-oriented programming can match when the resume supports that meaning.
- For skill match and missing lists, use the exact skill labels from the job's
  required_skills and preferred_skills arrays. Do not add skills to those lists.
- Base every conclusion on explicit evidence in the supplied resume and job.
- Set experience_requirement_met to null if no minimum experience is stated.
- Set education_requirement_satisfied to null if no education requirement is stated.
- relevant_experience must contain brief evidence from the resume, or be empty.
- Strengths must be specific to both the resume and this role.
- Quality issues must describe only supported resume omissions or weaknesses.
- Suggestions must be job-specific and grounded in existing resume content. Never
  tell the candidate to claim a skill, qualification, or experience they lack.
- Give a short verdict. Do not produce any score, percentage, or numerical rating.
"""
    user_prompt = (
        "JOB DESCRIPTION (structured):\n"
        f"{job.model_dump_json(indent=2)}\n\n"
        "RESUME (structured):\n"
        f"{resume.model_dump_json(indent=2)}"
    )
    data = _request_json(client, system_prompt, user_prompt)
    try:
        return MatchAnalysis.model_validate(data)
    except Exception as error:
        raise ValueError(f"Could not validate the match analysis: {error}") from error


def _canonical_matches(matches: list[str], job_skills: list[str]) -> list[str]:
    labels = {skill.strip().casefold(): skill for skill in job_skills}
    matched = {labels[skill.strip().casefold()] for skill in matches if skill.strip().casefold() in labels}
    return [skill for skill in job_skills if skill in matched]


def final_score(job: JobDescription, analysis: MatchAnalysis) -> tuple[float, ScoreBreakdown]:
    matching_required = _canonical_matches(analysis.matching_required_skills, job.required_skills)
    matching_preferred = _canonical_matches(analysis.matching_preferred_skills, job.preferred_skills)
    required_points = (
        len(matching_required) / len(job.required_skills) * 50 if job.required_skills else 50
    )
    preferred_points = (
        len(matching_preferred) / len(job.preferred_skills) * 15 if job.preferred_skills else 15
    )
    experience_points = (
        15 if job.minimum_experience is None or analysis.experience_requirement_met else 0
    )
    education_points = (
        10 if not job.education_requirements or analysis.education_requirement_satisfied else 0
    )
    relevant_points = 10 if analysis.relevant_experience else 0

    breakdown = ScoreBreakdown(
        required_skills=round(required_points, 2),
        preferred_skills=round(preferred_points, 2),
        experience=experience_points,
        education=education_points,
        relevant_experience=relevant_points,
    )
    total = round(sum(breakdown.model_dump().values()), 2)
    return total, breakdown


def read_pdf(file_path: Path) -> str:
    reader = PdfReader(file_path)
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


def read_docx(file_path: Path) -> str:
    document = Document(file_path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    return "\n".join(parts).strip()


def read_resume(file_path: Path) -> str:
    suffix = file_path.suffix.casefold()
    if suffix == ".pdf":
        text = read_pdf(file_path)
    elif suffix == ".docx":
        text = read_docx(file_path)
    else:
        raise ValueError(f"Unsupported resume file type: {file_path.suffix or '(no extension)'}")
    if not text:
        raise ValueError("No readable text was found in the resume.")
    return text


def _resume_information(resume: Resume) -> ResumeInformation:
    return ResumeInformation(
        name=bool(resume.name),
        email=bool(resume.email),
        phone=bool(resume.phone),
        skills=bool(resume.skills),
        experience=bool(resume.experiences),
        education=bool(resume.education),
        projects=bool(resume.projects),
        certifications=bool(resume.certifications),
    )


def process_resume(file_path: Path, job: JobDescription, client: Groq) -> CandidateResult:
    resume = parse_resume(read_resume(file_path), client)
    analysis = analyze_match(job, resume, client)
    matching_required = _canonical_matches(analysis.matching_required_skills, job.required_skills)
    matching_preferred = _canonical_matches(analysis.matching_preferred_skills, job.preferred_skills)
    missing_required = [skill for skill in job.required_skills if skill not in matching_required]
    missing_preferred = [skill for skill in job.preferred_skills if skill not in matching_preferred]
    score, breakdown = final_score(job, analysis)
    return CandidateResult(
        candidate_name=resume.name or file_path.stem,
        final_score=score,
        score_breakdown=breakdown,
        matching_required_skills=matching_required,
        missing_required_skills=missing_required,
        matching_preferred_skills=matching_preferred,
        missing_preferred_skills=missing_preferred,
        experience_requirement_met=analysis.experience_requirement_met,
        education_requirement_satisfied=analysis.education_requirement_satisfied,
        relevant_experience=analysis.relevant_experience,
        strengths=analysis.strengths,
        improvement_suggestions=analysis.improvement_suggestions,
        resume_information=_resume_information(resume),
        resume_quality_issues=analysis.quality_issues,
        verdict=analysis.verdict,
    )


def _status(value: bool | None) -> str:
    return "Yes" if value is True else "No" if value is False else "N/A"


def display_results(results: list[CandidateResult]) -> None:
    name_width = max(9, *(len(result.candidate_name) for result in results))
    header = (
        f"{'Candidate':<{name_width}}  {'Score':>7}  "
        f"{'Required Skills':>15}  {'Experience':>10}"
    )
    print("\n" + header)
    print("-" * len(header))
    for result in results:
        required_count = len(result.matching_required_skills)
        total_required = required_count + len(result.missing_required_skills)
        print(
            f"{result.candidate_name:<{name_width}}  {result.final_score:>7.2f}  "
            f"{required_count:>6}/{total_required:<8}  "
            f"{_status(result.experience_requirement_met):>10}"
        )

    for result in results:
        print(f"\n{result.candidate_name} | Overall Score: {result.final_score:.2f}/100")
        breakdown = result.score_breakdown
        print(f"  Required Skills:       {breakdown.required_skills:>5.2f}/50")
        print(f"  Preferred Skills:      {breakdown.preferred_skills:>5.2f}/15")
        print(f"  Experience:            {breakdown.experience:>5.2f}/15")
        print(f"  Education:              {breakdown.education:>5.2f}/10")
        print(f"  Relevant Experience:   {breakdown.relevant_experience:>5.2f}/10")
        print(f"  Matching required: {', '.join(result.matching_required_skills) or 'None'}")
        print(f"  Missing required:  {', '.join(result.missing_required_skills) or 'None'}")
        print(f"  Matching preferred: {', '.join(result.matching_preferred_skills) or 'None'}")
        print(f"  Missing preferred:  {', '.join(result.missing_preferred_skills) or 'None'}")
        print(f"  Experience requirement met: {_status(result.experience_requirement_met)}")
        print(f"  Education requirement met:  {_status(result.education_requirement_satisfied)}")
        print(f"  Strengths: {', '.join(result.strengths) or 'None identified'}")
        print(f"  Verdict: {result.verdict or 'No verdict returned.'}")
        print("  Resume information:")
        for field, present in result.resume_information.model_dump().items():
            print(f"    {field.title()}: {'Present' if present else 'Missing'}")
        print(f"  Quality issues: {', '.join(result.resume_quality_issues) or 'None identified'}")
        print(f"  Improvement suggestions: {', '.join(result.improvement_suggestions) or 'None'}")


def export_json(results: list[CandidateResult], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps([result.model_dump() for result in results], indent=2),
        encoding="utf-8",
    )


def export_csv(results: list[CandidateResult], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    columns = list(CandidateResult.model_fields)
    with output_path.open("w", newline="", encoding="utf-8") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=columns)
        writer.writeheader()
        for result in results:
            row = result.model_dump()
            writer.writerow({
                key: json.dumps(value) if isinstance(value, (dict, list)) else value
                for key, value in row.items()
            })


def _read_job_description(path: Path | None) -> str:
    if path is not None:
        text = path.read_text(encoding="utf-8").strip()
    else:
        print("Paste the job description, then finish input with Ctrl+Z and Enter:")
        text = sys.stdin.read().strip()
    if not text:
        raise ValueError("The job description is empty.")
    return text


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare resumes with a job description.")
    parser.add_argument("--job-description-file", type=Path, help="Text file containing the job description")
    parser.add_argument("--resumes-dir", type=Path, default=DEFAULT_RESUMES_DIR)
    parser.add_argument("--json-output", type=Path, help="Export results to this JSON file")
    parser.add_argument("--csv-output", type=Path, help="Export results to this CSV file")
    args = parser.parse_args()

    load_dotenv()
    api_key = os.getenv("GROQ_API_KEY")
    if not api_key:
        parser.error("GROQ_API_KEY is not set. Add it to your environment or .env file.")
    if not args.resumes_dir.is_dir():
        parser.error(f"Resume directory does not exist: {args.resumes_dir}")

    try:
        job_text = _read_job_description(args.job_description_file)
        client = Groq(api_key=api_key)
        job = parse_job_description(job_text, client)
    except Exception as error:
        parser.exit(1, f"Could not prepare job description: {error}\n")

    results: list[CandidateResult] = []
    for file_path in sorted(args.resumes_dir.iterdir()):
        if not file_path.is_file():
            continue
        print(f"\nProcessing: {file_path.name}")
        try:
            result = process_resume(file_path, job, client)
        except Exception as error:
            print(f"  Error: {error}. Skipping this resume.", file=sys.stderr)
            continue
        results.append(result)

    results.sort(key=lambda result: result.final_score, reverse=True)
    if results:
        display_results(results)
    else:
        print("No resumes were successfully processed.")

    if args.json_output:
        export_json(results, args.json_output)
        print(f"\nJSON results saved to {args.json_output}")
    if args.csv_output:
        export_csv(results, args.csv_output)
        print(f"CSV results saved to {args.csv_output}")


__all__ = [
    "main",
    "MODEL",
    "JobDescription",
    "Resume",
    "MatchAnalysis",
    "ScoreBreakdown",
    "CandidateResult",
    "parse_job_description",
    "parse_resume",
    "analyze_match",
    "final_score",
    "process_resume",
    "read_resume",
    "_canonical_matches",
    "_request_json",
]
